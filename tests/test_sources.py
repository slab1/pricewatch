"""Tests for source config, fetchers, the NBS importer and demo seeding.
No network: fetchers get fake responses or html_overrides; every DB write goes
to a temp file via monkeypatched db.DB_PATH."""
import json
import os
import sys
from datetime import date

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

import db  # noqa: E402
import import_nbs  # noqa: E402
import seed_demo  # noqa: E402
import sources  # noqa: E402

CONFIG = os.path.join(_ROOT, "sources.json")


@pytest.fixture
def tmpdb(monkeypatch, tmp_path):
    path = str(tmp_path / "t.db")
    monkeypatch.setattr(db, "DB_PATH", path)
    return path


# ---- sources.json integrity -------------------------------------------------

def test_sources_json_is_valid_and_complete():
    entries = json.loads(open(CONFIG).read())
    assert isinstance(entries, list) and entries
    for e in entries:
        for key in ("item", "source", "unit", "type", "url", "min", "max", "enabled"):
            assert key in e, f"{e.get('item')}/{e.get('source')} missing {key}"
        assert e["type"] in sources.FETCHERS, f"unknown type {e['type']}"
        assert e["min"] < e["max"]
        if not e["enabled"]:
            assert "_note" in e and e["_note"].strip(), \
                f"disabled source {e['item']}/{e['source']} must explain why in _note"


def test_every_item_has_three_documented_candidates():
    entries = json.loads(open(CONFIG).read())
    by_item = {}
    for e in entries:
        by_item.setdefault(e["item"], []).append(e)
    assert len(by_item) >= 10, f"expected >=10 items, got {len(by_item)}"
    for item, rows in by_item.items():
        assert len(rows) >= 3, f"{item} has only {len(rows)} candidate sources"
        keys = [(r["item"], r["source"]) for r in rows]
        assert len(keys) == len(set(keys)), f"{item} has duplicate source rows"


def test_enabled_sources_have_live_provenance():
    """Only verified sources may be enabled; each enabled one carries a URL and
    a note if any. Guards against flipping an unverified entry on."""
    entries = json.loads(open(CONFIG).read())
    enabled = [e for e in entries if e["enabled"]]
    assert enabled, "at least one source must be enabled"
    for e in enabled:
        assert e["url"].startswith("http")
        if "_note" in e:
            assert "unverified" not in e["_note"].lower(), \
                f"{e['item']}/{e['source']} enabled but note calls it unverified"


# ---- fetchers ---------------------------------------------------------------

@pytest.mark.parametrize("text,expected", [
    ("\u20a61,234.50 per litre", 1234.5),
    ("1331", 1331.0),
    ("1,331.01", 1331.01),
    ("NGN 875.5/kg", 875.5),
])
def test_parse_number(text, expected):
    assert sources.parse_number(text) == expected


def test_parse_number_raises_without_digits():
    with pytest.raises(ValueError, match="no number"):
        sources.parse_number("out of stock")


def test_check_range_bounds():
    assert sources.check_range(500, {"min": 100, "max": 1000}) == 500
    with pytest.raises(ValueError, match="outside sanity range"):
        sources.check_range(50, {"min": 100, "max": 1000})
    with pytest.raises(ValueError):
        sources.check_range(5000, {"min": 100, "max": 1000})


def test_build_headers_expands_env(monkeypatch):
    monkeypatch.setenv("PW_TEST_TOKEN", "secret123")
    h = sources.build_headers({"headers": {"X-Token": "Bearer ${PW_TEST_TOKEN}"}})
    assert h["X-Token"] == "Bearer secret123"


def test_build_headers_rejects_unset_env(monkeypatch):
    monkeypatch.delenv("PW_TEST_TOKEN", raising=False)
    with pytest.raises(ValueError, match="environment variable"):
        sources.build_headers({"headers": {"X-Token": "${PW_TEST_TOKEN}"}})


def test_html_css_fetch_with_override():
    cfg = {"selector": ".price", "min": 1, "max": 10000}
    assert sources.fetch_html_css(cfg, "<div class='price'>\u20a61,331</div>") == 1331.0
    with pytest.raises(ValueError, match="matched nothing"):
        sources.fetch_html_css(cfg, "<div class='nope'>x</div>")


class _FakeResp:
    def __init__(self, payload, status=200):
        self._payload, self.status_code = payload, status

    def raise_for_status(self):
        if self.status_code >= 400:
            import requests
            raise requests.HTTPError(str(self.status_code))

    def json(self):
        return self._payload


def test_json_api_fetch(monkeypatch):
    monkeypatch.setattr(sources.requests, "get",
                        lambda url, **kw: _FakeResp({"rates": {"NGN": 1331.01}}))
    cfg = {"url": "http://x", "path": ["rates", "NGN"], "min": 100, "max": 10000}
    assert sources.fetch_json_api(cfg) == 1331.01


def test_json_api_http_error_propagates(monkeypatch):
    monkeypatch.setattr(sources.requests, "get", lambda url, **kw: _FakeResp({}, 503))
    with pytest.raises(Exception):
        sources.fetch_json_api({"url": "http://x", "path": ["a"], "min": 1, "max": 9})


# ---- month / filename parsing ----------------------------------------------

@pytest.mark.parametrize("text,expected", [
    ("April 2026", "2026-04-15"),
    ("Apr-2026", "2026-04-15"),
    ("2026-06", "2026-06-15"),
    ("June, 2025", "2025-06-15"),
])
def test_parse_month(text, expected):
    assert import_nbs.parse_month(text) == expected


def test_month_from_filename():
    assert import_nbs.month_from_filename("/x/petrol_may_2026.xlsx") == "2026-05-15"
    assert import_nbs.month_from_filename("/x/nodate.xlsx") is None


def test_match_item_avoids_substring_traps():
    # 'rice' must not match inside 'price', nor 'ago' inside 'Lagos'
    assert import_nbs.match_item("Price Per Kg (NGN)") is None
    assert import_nbs.match_item("Lagos - Abuja urban bus fare")[0] == "bus_fare"
    assert import_nbs.match_item("Rice (imported)")[0] == "rice"
    assert import_nbs.match_item("Tomatoes (red)")[0] == "tomatoes"
    assert import_nbs.match_item("Beans (brown)")[0] == "beans"


# ---- NBS importer -----------------------------------------------------------

def _petrol_workbook(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Petrol (PMS) Price Watch"
    ws.append(["Petrol (PMS) Price Watch", None, None])
    ws.append(["Month of Survey", "PMS Price Per Litre (N)", "National Average"])
    ws.append(["April 2026", 869.0, 869.0])
    ws.append(["May 2026", 875.5, 875.5])
    wb.save(path)
    return path


def _food_workbook(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Selected Food Prices"
    ws.append(["Item", "Unit", "Price Per Kg (NGN)", "National Average (NGN)"])
    ws.append(["Rice (imported)", "Kg", 1900, 1900])
    ws.append(["Beans (brown)", "Kg", 1750, 1750])
    ws.append(["Yam", "Kg", 900, 900])
    wb.save(path)
    return path


def test_import_single_commodity_report(tmp_path, tmpdb):
    f = _petrol_workbook(str(tmp_path / "petrol_april_2026.xlsx"))
    found, rejected = import_nbs.import_file(f)
    rows = {(r[0], r[3]): r[1] for r in found}
    assert rows[("petrol_pms", "2026-04-15")] == 869.0
    assert rows[("petrol_pms", "2026-05-15")] == 875.5
    with db.connect() as c:
        stored = c.execute("SELECT item, value, collected_date, source FROM prices ORDER BY collected_date").fetchall()
    assert [s["collected_date"] for s in stored] == ["2026-04-15", "2026-05-15"]
    assert all(s["source"] == "nbs" for s in stored)


def test_import_multi_row_food_report(tmp_path, tmpdb):
    f = _food_workbook(str(tmp_path / "food_april_2026.xlsx"))
    found, rejected = import_nbs.import_file(f)
    got = {r[0]: r[1] for r in found}
    assert got["rice"] == 1900.0 and got["beans"] == 1750.0
    assert "yam" not in got  # unmatched rows are skipped, not guessed


def test_import_dry_run_writes_nothing(tmp_path, tmpdb):
    f = _petrol_workbook(str(tmp_path / "petrol_april_2026.xlsx"))
    found, _ = import_nbs.import_file(f, dry_run=True)
    assert len(found) == 2
    with db.connect() as c:
        assert c.execute("SELECT COUNT(*) FROM prices").fetchone()[0] == 0


def test_import_rejects_out_of_range(tmp_path, tmpdb):
    import openpyxl
    path = str(tmp_path / "petrol_april_2026.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Petrol (PMS) Price Watch"
    ws.append(["Month of Survey", "PMS Price Per Litre (N)", "National Average"])
    ws.append(["April 2026", 999999, 999999])
    wb.save(path)
    found, rejected = import_nbs.import_file(path)
    assert found == []
    assert any("outside sanity" in r for r in rejected)


def test_import_item_filter(tmp_path, tmpdb):
    f = _food_workbook(str(tmp_path / "food_april_2026.xlsx"))
    found, _ = import_nbs.import_file(f, item_filter="rice")
    assert [r[0] for r in found] == ["rice"]


def test_import_is_idempotent_per_month(tmp_path, tmpdb):
    f = _petrol_workbook(str(tmp_path / "petrol_april_2026.xlsx"))
    import_nbs.import_file(f)
    import_nbs.import_file(f)
    with db.connect() as c:
        assert c.execute("SELECT COUNT(*) FROM prices").fetchone()[0] == 2


# ---- demo seeding -----------------------------------------------------------

def test_seed_demo_creates_labeled_rows(tmpdb):
    n = seed_demo.seed(5, today=date(2026, 10, 9))
    assert n == len(seed_demo.ITEMS) * 5
    with db.connect() as c:
        rows = c.execute("SELECT DISTINCT source FROM prices").fetchall()
        assert [r["source"] for r in rows] == ["demo"]
        assert c.execute("SELECT MIN(value) FROM prices").fetchone()[0] > 0


def test_seed_demo_is_idempotent(tmpdb):
    seed_demo.seed(7, today=date(2026, 10, 9))
    seed_demo.seed(7, today=date(2026, 10, 9))
    with db.connect() as c:
        assert c.execute("SELECT COUNT(*) FROM prices").fetchone()[0] == len(seed_demo.ITEMS) * 7


# ---- db.save_price date parameter (backward compatible) ---------------------

def test_save_price_defaults_to_today(tmpdb):
    db.save_price("rice", "nbs", 1900.0, "NGN per kg")
    with db.connect() as c:
        row = c.execute("SELECT collected_date FROM prices").fetchone()
    assert row["collected_date"] == date.today().isoformat()


def test_save_price_backfills_and_upserts(tmpdb):
    db.save_price("rice", "nbs", 1900.0, "NGN per kg", when="2026-04-15")
    db.save_price("rice", "nbs", 1950.0, "NGN per kg", when="2026-04-15")
    with db.connect() as c:
        rows = c.execute("SELECT value, collected_date FROM prices").fetchall()
    assert len(rows) == 1
    assert rows[0]["value"] == 1950.0 and rows[0]["collected_date"] == "2026-04-15"
